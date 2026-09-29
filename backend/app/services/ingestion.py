import csv
import io
import zipfile
from pathlib import Path
from fastapi import HTTPException
from sqlalchemy import select
from app.models import Material, Dataset, Evidence
from app.audit import audit

MAX_UPLOAD = 10 * 1024 * 1024


def parse_file(filename, data):
    extension = Path(filename).suffix.lower()
    if len(data) > MAX_UPLOAD:
        raise HTTPException(413, "File exceeds the 10 MB limit.")
    try:
        if extension == ".csv":
            result = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
        elif extension == ".xlsx":
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(x.file_size for x in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValueError("Expanded workbook is too large")
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            iterator = workbook.active.iter_rows(values_only=True)
            headers = [str(x or "").strip() for x in next(iterator)]
            result = [dict(zip(headers, [str(x) if x is not None else "" for x in row])) for row in iterator]
            workbook.close()
        elif extension in (".pdf", ".txt"):
            if extension == ".pdf":
                from pypdf import PdfReader
                pages = [(i + 1, p.extract_text() or "") for i, p in enumerate(PdfReader(io.BytesIO(data)).pages)]
            else:
                pages = [(None, data.decode("utf-8-sig"))]
            result = [{"legacy_material_code": "", "original_description": line.strip(), "category": "", "page_number": page} for page, text in pages for line in text.splitlines() if line.strip()]
        else:
            raise HTTPException(415, "Use a CSV, XLSX, text PDF or TXT file.")
        if not result or len(result) > 10000:
            raise ValueError("Empty file or too many rows")
        if any(len(str(v)) > 10000 for row in result for v in row.values()):
            raise ValueError("Cell exceeds maximum length")
        return result
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, "Unable to process the uploaded file. Use UTF-8 CSV, XLSX or a text-based document (maximum 10,000 rows).")


def validate(db, dataset, mapping):
    if not all(mapping.get(k) for k in ("legacy_material_code", "original_description", "category")):
        raise HTTPException(422, "Column mapping is incomplete. Map code, description and category.")
    existing = set(db.scalars(select(Material.legacy_material_code).where(Material.organization_id == dataset.organization_id)))
    seen, report = set(), []
    for index, row in enumerate(dataset.source_rows, 2):
        values = {k: str(row.get(v, "") or "").strip() for k, v in mapping.items()}
        errors = []
        code = values.get("legacy_material_code", "")
        if not code:
            errors.append("Missing material code")
        elif len(code) > 120:
            errors.append("Material code exceeds 120 characters")
        elif code in seen or code in existing:
            errors.append("Duplicate source code")
        seen.add(code)
        if len(values.get("original_description", "")) < 3:
            errors.append("Description is missing or too short")
        if not values.get("category") or len(values["category"]) > 100:
            errors.append("Category is missing or too long")
        if values.get("unit") and values["unit"].lower() not in ("ea", "each", "nos", "kg", "m", "mm", "l", "set", "pc", "pcs"):
            errors.append("Invalid unit")
        report.append({"row": index, "values": values, "errors": errors, "page_number": row.get("page_number")})
    dataset.column_mapping = mapping
    dataset.validation_report = report
    dataset.valid_rows = sum(not row["errors"] for row in report)
    dataset.invalid_rows = len(report) - dataset.valid_rows
    dataset.status = "VALIDATED"
    return report


def import_dataset(db, dataset, user):
    if dataset.status == "IMPORTED":
        raise HTTPException(409, "This dataset has already been imported.")
    if dataset.status != "VALIDATED":
        raise HTTPException(409, "Validate column mapping before importing.")
    validate(db, dataset, dataset.column_mapping)
    for row in dataset.validation_report:
        if row["errors"]:
            continue
        values = row["values"]
        description = values["original_description"]
        extra = [f"{k}: {v}" for k, v in values.items() if k not in ("original_description", "legacy_material_code", "category") and v]
        if extra:
            description += " | " + " | ".join(extra)
        material = Material(organization_id=dataset.organization_id, dataset_id=dataset.id, legacy_material_code=values["legacy_material_code"], original_description=description, category=values["category"])
        db.add(material)
        audit(db, user, "MATERIAL_UPLOADED", material, new=values, reason=f"Source row {row['row']}; synthetic or user-supplied data")
        db.add(Evidence(material_id=material.id, source_type=dataset.source_type, source_file=dataset.filename, row_number=row["row"], page_number=row.get("page_number"), source_text=description, extracted_value=None, evidence_type="SOURCE", extraction_method="RULE"))
    dataset.status = "IMPORTED"
    audit(db, user, "DATASET_IMPORTED", dataset, new={"imported": dataset.valid_rows, "skipped": dataset.invalid_rows}, reason="Only validated rows imported")
    return {"imported": dataset.valid_rows, "skipped": dataset.invalid_rows}
